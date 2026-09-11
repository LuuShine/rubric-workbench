import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, List, Optional


SUPPORTED_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".pdf", ".docx"}
MAX_FILES = 6
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_CHARS_PER_FILE = 8_000
MAX_TOTAL_CHARS = 20_000


@dataclass(frozen=True)
class ReferenceSummary:
    name: str
    status: str
    chars: int
    message: str


@dataclass(frozen=True)
class ReferenceBundle:
    context: str
    summaries: List[ReferenceSummary]

    @property
    def loaded_count(self) -> int:
        return sum(item.status == "loaded" for item in self.summaries)

    @property
    def total_chars(self) -> int:
        return sum(item.chars for item in self.summaries if item.status == "loaded")


def _as_list(files: Any) -> List[Any]:
    if files is None:
        return []
    if isinstance(files, (str, Path, dict)):
        return [files]
    if isinstance(files, Iterable):
        return list(files)
    return [files]


def _resolve_path(uploaded_file: Any) -> Optional[Path]:
    if isinstance(uploaded_file, (str, Path)):
        return Path(uploaded_file)
    if isinstance(uploaded_file, dict):
        value = uploaded_file.get("path") or uploaded_file.get("name")
        return Path(value) if value else None
    value = getattr(uploaded_file, "name", None)
    return Path(value) if value else None


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_json(path: Path) -> str:
    value = json.loads(_decode_text(path.read_bytes()))
    return json.dumps(value, ensure_ascii=False, indent=2)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[第 {index} 页]\n{text}")
    if not pages:
        raise ValueError("PDF 中没有可提取文本，扫描件需先进行 OCR。")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    text = "\n".join(block for block in blocks if block)
    if not text:
        raise ValueError("DOCX 中没有可提取文本。")
    return text


def _read_reference(path: Path) -> str:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持 {extension or '无扩展名'} 文件。")
    if not path.exists() or not path.is_file():
        raise ValueError("上传文件不存在或已失效。")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("文件超过 10 MB 限制。")

    if extension == ".json":
        text = _read_json(path)
    elif extension == ".pdf":
        text = _read_pdf(path)
    elif extension == ".docx":
        text = _read_docx(path)
    else:
        text = _decode_text(path.read_bytes())

    text = text.replace("\x00", "").strip()
    if not text:
        raise ValueError("文件中没有可用文本。")
    return text


def extract_reference_files(files: Any) -> ReferenceBundle:
    uploaded_files = _as_list(files)
    summaries: List[ReferenceSummary] = []
    context_blocks = []
    remaining_chars = MAX_TOTAL_CHARS

    for index, uploaded_file in enumerate(uploaded_files):
        path = _resolve_path(uploaded_file)
        name = path.name if path else f"文件 {index + 1}"
        if index >= MAX_FILES:
            summaries.append(
                ReferenceSummary(name, "skipped", 0, f"最多读取 {MAX_FILES} 个文件。")
            )
            continue
        if path is None:
            summaries.append(
                ReferenceSummary(name, "error", 0, "无法识别上传文件路径。")
            )
            continue

        try:
            text = _read_reference(path)
        except Exception as error:
            # Uploaded documents are untrusted and parser libraries expose
            # format-specific exception types.
            summaries.append(ReferenceSummary(name, "error", 0, str(error)))
            continue

        allowed_chars = min(MAX_CHARS_PER_FILE, remaining_chars)
        if allowed_chars <= 0:
            summaries.append(
                ReferenceSummary(name, "skipped", 0, "参考资料总长度已达到限制。")
            )
            continue

        original_chars = len(text)
        text = text[:allowed_chars]
        remaining_chars -= len(text)
        truncated = len(text) < original_chars
        message = (
            f"已读取 {len(text)} 字符，超出部分已截断。"
            if truncated
            else f"已读取 {len(text)} 字符。"
        )
        summaries.append(ReferenceSummary(name, "loaded", len(text), message))
        context_blocks.append(f"### 文件：{name}\n{text}")

    return ReferenceBundle(
        context="\n\n".join(context_blocks),
        summaries=summaries,
    )
