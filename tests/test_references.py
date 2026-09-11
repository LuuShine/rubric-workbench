from docx import Document

from rubric_tool.references import (
    MAX_CHARS_PER_FILE,
    extract_reference_files,
)


def test_extracts_multiple_text_and_json_files(tmp_path) -> None:
    policy = tmp_path / "policy.md"
    policy.write_text("退款审核需要检查订单状态。", encoding="utf-8")
    cases = tmp_path / "cases.json"
    cases.write_text('{"high_risk": ["重复退款"]}', encoding="utf-8")

    bundle = extract_reference_files([str(policy), str(cases)])

    assert bundle.loaded_count == 2
    assert "退款审核需要检查订单状态" in bundle.context
    assert '"high_risk"' in bundle.context
    assert "重复退款" in bundle.context


def test_truncates_long_reference_file(tmp_path) -> None:
    reference = tmp_path / "long.txt"
    reference.write_text("a" * (MAX_CHARS_PER_FILE + 200), encoding="utf-8")

    bundle = extract_reference_files(str(reference))

    assert bundle.loaded_count == 1
    assert bundle.summaries[0].chars == MAX_CHARS_PER_FILE
    assert "已截断" in bundle.summaries[0].message


def test_reports_unsupported_file_without_loading_it(tmp_path) -> None:
    reference = tmp_path / "archive.zip"
    reference.write_bytes(b"not-a-real-archive")

    bundle = extract_reference_files(str(reference))

    assert bundle.loaded_count == 0
    assert bundle.summaries[0].status == "error"
    assert "不支持" in bundle.summaries[0].message


def test_extracts_docx_paragraphs_and_table_cells(tmp_path) -> None:
    reference = tmp_path / "policy.docx"
    document = Document()
    document.add_paragraph("退款规则")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "到账方式"
    table.cell(0, 1).text = "原路退回"
    document.save(reference)

    bundle = extract_reference_files(str(reference))

    assert bundle.loaded_count == 1
    assert "退款规则" in bundle.context
    assert "到账方式 | 原路退回" in bundle.context
